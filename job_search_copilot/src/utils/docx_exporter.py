"""ATS-oriented DOCX resume export."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt


def _add_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(6)


def _add_block(doc: Document, title: str, body: str | None) -> None:
    if not body or not str(body).strip():
        return
    _add_heading(doc, title)
    for line in str(body).strip().split("\n"):
        if line.strip():
            doc.add_paragraph(line.strip())


def build_resume_docx(
    *,
    output_path: Path,
    user: dict[str, Any],
    resume: dict[str, Any],
    profile: dict[str, Any] | None = None,
) -> Path:
    """Write a simple, ATS-friendly resume DOCX."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    name = (user.get("name") or "Candidate").strip()
    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r = title.add_run(name)
    r.bold = True
    r.font.size = Pt(16)

    contact_bits = []
    if user.get("email"):
        contact_bits.append(str(user["email"]))
    if user.get("phone"):
        contact_bits.append(str(user["phone"]))
    if user.get("current_location"):
        contact_bits.append(str(user["current_location"]))
    if contact_bits:
        cp = doc.add_paragraph(" | ".join(contact_bits))
        cp.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        for run in cp.runs:
            run.font.size = Pt(10)

    headline = (resume.get("resume_headline") or "").strip()
    if headline:
        hp = doc.add_paragraph(headline)
        hp.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        for run in hp.runs:
            run.italic = True
            run.font.size = Pt(11)

    doc.add_paragraph()

    _add_block(doc, "Professional Summary", resume.get("professional_summary"))
    _add_block(doc, "Skills", resume.get("tailored_skills"))
    _add_block(doc, "Experience", resume.get("tailored_experience"))

    prof = profile or {}
    _add_block(doc, "Projects", prof.get("projects"))
    _add_block(doc, "Achievements", prof.get("achievements"))

    final_text = (resume.get("final_resume_text") or "").strip()
    if final_text:
        _add_block(doc, "Consolidated Resume", final_text)

    _add_block(doc, "Education", prof.get("education"))
    _add_block(doc, "Certifications", prof.get("certifications"))

    doc.save(str(output_path))
    return output_path
